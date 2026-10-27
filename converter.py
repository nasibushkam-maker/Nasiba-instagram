import csv
import html
import os
import shutil
import subprocess
import uuid
from pathlib import Path
from typing import List, Tuple, Optional
from PIL import Image

CONVERT_DIR = Path("converted_files")
CONVERT_DIR.mkdir(parents=True, exist_ok=True)

# Форматлар категориялари
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tiff", ".tif", ".ico"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".ogg", ".m4a", ".aac", ".flac", ".wma", ".opus"}
VIDEO_EXTENSIONS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".wmv", ".flv", ".3gp"}
EXCEL_EXTENSIONS = {".xlsx", ".xls", ".csv"}
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt"}


def get_file_category(extension: str) -> Optional[str]:
    ext = extension.lower()
    if ext in IMAGE_EXTENSIONS:
        return "image"
    if ext in AUDIO_EXTENSIONS:
        return "audio"
    if ext in VIDEO_EXTENSIONS:
        return "video"
    if ext in EXCEL_EXTENSIONS:
        return "excel"
    if ext in DOCUMENT_EXTENSIONS:
        return "document"
    return None


def get_available_target_formats(extension: str) -> List[Tuple[str, str]]:
    """
    Файл кенгайтмасига қараб, мос келадиган мақсадли форматларни қайтаради.
    Қайтариладиган қиймат: [("Тугма матни", "формат_код")]
    """
    ext = extension.lower()
    cat = get_file_category(ext)

    if cat == "excel":
        if ext == ".xlsx":
            return [
                ("📑 PDF (PDF ҳужжат)", "pdf"),
                ("📊 CSV (Матнли жадвал)", "csv"),
                ("🌐 HTML (Веб саҳифа)", "html"),
            ]
        elif ext == ".xls":
            return [
                ("📗 XLSX (Янги Excel)", "xlsx"),
                ("📑 PDF (PDF ҳужжат)", "pdf"),
                ("📊 CSV (Матнли жадвал)", "csv"),
            ]
        elif ext == ".csv":
            return [
                ("📗 XLSX (Excel формати)", "xlsx"),
                ("📑 PDF (PDF ҳужжат)", "pdf"),
                ("🌐 HTML (Веб саҳифа)", "html"),
            ]

    if cat == "image":
        formats = [
            ("🖼 PNG", "png"),
            ("🖼 JPG", "jpg"),
            ("🌐 WEBP", "webp"),
            ("📑 PDF", "pdf"),
            ("⭐ ICO", "ico"),
            ("🎨 BMP", "bmp"),
        ]
        current = ext.lstrip(".")
        if current == "jpeg":
            current = "jpg"
        return [f for f in formats if f[1] != current]

    if cat == "audio":
        formats = [
            ("🎵 MP3", "mp3"),
            ("🔊 WAV", "wav"),
            ("🎙 OGG", "ogg"),
            ("🎶 M4A", "m4a"),
            ("🎧 AAC", "aac"),
        ]
        current = ext.lstrip(".")
        return [f for f in formats if f[1] != current]

    if cat == "video":
        formats = [
            ("🎵 MP3 (Аудио олиш)", "mp3"),
            ("🎞 GIF (Анимация)", "gif"),
            ("🎬 MP4", "mp4"),
            ("🎥 AVI", "avi"),
            ("📼 MKV", "mkv"),
        ]
        current = ext.lstrip(".")
        return [f for f in formats if f[1] != current]

    if cat == "document":
        if ext == ".pdf":
            return [
                ("📝 DOCX (Word)", "docx"),
                ("📄 TXT (Матн)", "txt"),
                ("🖼 PNG (Расмлар)", "png_images"),
            ]
        elif ext == ".docx":
            return [
                ("📑 PDF", "pdf"),
            ]
        elif ext == ".txt":
            return [
                ("📑 PDF", "pdf"),
                ("📝 DOCX", "docx"),
            ]

    return []


# ==========================================
# EXCEL / CSV ЖАДВАЛЛАРНИ ҚАЙТА ИШЛАШ
# ==========================================
def read_table_rows(input_path: str) -> List[List[str]]:
    """XLSX, XLS ёки CSV дан қаторларни ўқиб, матнлар рўйхати сифатида қайтаради"""
    _, ext = os.path.splitext(input_path)
    ext = ext.lower()
    rows_data: List[List[str]] = []

    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(input_path, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if any(c is not None and str(c).strip() != "" for c in row):
                rows_data.append([str(c) if c is not None else "" for c in row])

    elif ext == ".xls":
        import xlrd
        wb = xlrd.open_workbook(input_path)
        ws = wb.sheet_by_index(0)
        for r in range(ws.nrows):
            row = ws.row_values(r)
            if any(c is not None and str(c).strip() != "" for c in row):
                rows_data.append([str(c) if c is not None else "" for c in row])

    elif ext == ".csv":
        for encoding in ("utf-8-sig", "utf-8", "windows-1251", "cp1252"):
            try:
                with open(input_path, "r", encoding=encoding, errors="replace") as f:
                    sample = f.read(2048)
                    f.seek(0)
                    try:
                        dialect = csv.Sniffer().sniff(sample)
                        reader = csv.reader(f, dialect)
                    except Exception:
                        reader = csv.reader(f)
                    for row in reader:
                        if any(c.strip() != "" for c in row):
                            rows_data.append([str(c) for c in row])
                break
            except Exception:
                continue

    return rows_data


def find_browser() -> Optional[str]:
    """PDF яратиш учун Edge, Chrome ёки Chromium браузерини топади (Windows ва Linux учун)"""
    candidates = [
        # Linux / Docker / Render йўллари
        shutil.which("chromium"),
        shutil.which("chromium-browser"),
        shutil.which("google-chrome"),
        shutil.which("google-chrome-stable"),
        "/usr/bin/chromium",
        "/usr/bin/chromium-browser",
        "/usr/bin/google-chrome",
        # Windows йўллари
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        shutil.which("msedge"),
        shutil.which("chrome"),
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    ]
    for c in candidates:
        if c and os.path.exists(c):
            return c
    return None


def convert_table_to_pdf(input_path: str) -> str:
    """
    Excel / CSV жадвални PDF га айлантиради.
    Катта жадваллар (масалан, 147 қатор, 22 устун) учун ҳам кафолатли, музлаб қолмасдан ишлайди.
    """
    rows_data = read_table_rows(input_path)
    if not rows_data:
        raise ValueError("Жадвалда маълумот топилмади ёки файл бўш.")

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.pdf"
    output_path = str(CONVERT_DIR / output_filename)
    abs_output_path = os.path.abspath(output_path)

    browser = find_browser()
    if browser:
        temp_html = str(CONVERT_DIR / f"temp_{uuid.uuid4().hex[:8]}.html")
        abs_temp_html = os.path.abspath(temp_html)

        max_cols = max(len(r) for r in rows_data) if rows_data else 1
        # Устунлар сони кўп бўлса (масалан, 22 устун), форматни кенгроқ A3/A4 landscape га мослаш
        if max_cols > 15:
            page_size = "A3 landscape"
            font_size = "7pt"
            padding = "3px 4px"
        elif max_cols > 8:
            page_size = "A4 landscape"
            font_size = "7.5pt"
            padding = "4px 6px"
        else:
            page_size = "A4 portrait"
            font_size = "9pt"
            padding = "6px 8px"

        rows_html = []
        for row_idx, row in enumerate(rows_data):
            tag = "th" if row_idx == 0 else "td"
            cells = []
            for cell in row:
                escaped = html.escape(str(cell)) if cell is not None else ""
                cells.append(f"<{tag}>{escaped or '&nbsp;'}</{tag}>")
            rows_html.append(f"<tr>{''.join(cells)}</tr>")

        html_content = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
@page {{
    size: {page_size};
    margin: 8mm;
}}
body {{
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    margin: 0;
    padding: 0;
    color: #1a202c;
    background: #ffffff;
}}
table {{
    border-collapse: collapse;
    width: 100%;
    font-size: {font_size};
}}
tr {{
    page-break-inside: avoid;
}}
th, td {{
    border: 1px solid #cbd5e0;
    padding: {padding};
    text-align: left;
    vertical-align: top;
    word-break: break-word;
}}
th {{
    background-color: #1F4E79;
    color: #ffffff;
    font-weight: 600;
}}
tr:nth-child(even) td {{
    background-color: #f8fafc;
}}
</style>
</head>
<body>
<table>
{''.join(rows_html)}
</table>
</body>
</html>"""

        try:
            with open(abs_temp_html, "w", encoding="utf-8") as f:
                f.write(html_content)

            cmd = [
                browser,
                "--headless=new",
                "--no-sandbox",
                "--disable-setuid-sandbox",
                "--no-pdf-header-footer",
                "--disable-gpu",
                f"--print-to-pdf={abs_output_path}",
                abs_temp_html,
            ]
            subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=45)

            if os.path.exists(abs_output_path) and os.path.getsize(abs_output_path) > 0:
                return abs_output_path
        finally:
            if os.path.exists(abs_temp_html):
                try:
                    os.remove(abs_temp_html)
                except OSError:
                    pass

    # Браузер топилмаса ёки ишламаса, хавфсиз ReportLab варианти
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A3),
        leftMargin=15,
        rightMargin=15,
        topMargin=15,
        bottomMargin=15,
    )

    styles = getSampleStyleSheet()
    cell_style = ParagraphStyle("CStyle", parent=styles["Normal"], fontSize=7, leading=9)
    header_style = ParagraphStyle("HStyle", parent=styles["Normal"], fontSize=7, leading=9, textColor=colors.whitesmoke, fontName="Helvetica-Bold")

    max_cols = max(len(r) for r in rows_data)
    table_data = []

    for row_idx, row in enumerate(rows_data[:500]):
        padded = row + [""] * (max_cols - len(row))
        row_cells = []
        for cell in padded:
            # Узун матнларни чеклаш (хатолик бермаслиги учун)
            text_val = html.escape(str(cell)[:300]) if cell is not None else ""
            st = header_style if row_idx == 0 else cell_style
            row_cells.append(Paragraph(text_val or "&nbsp;", st))
        table_data.append(row_cells)

    t = Table(table_data, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F4E79")),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E0")),
    ]))

    doc.build([t])
    return output_path


def convert_excel_to_csv(input_path: str) -> str:
    """Excel ни CSV файлига айлантиради"""
    rows_data = read_table_rows(input_path)
    output_filename = f"converted_{uuid.uuid4().hex[:8]}.csv"
    output_path = str(CONVERT_DIR / output_filename)

    with open(output_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerows(rows_data)

    return output_path


def convert_excel_to_html(input_path: str) -> str:
    """Excel ни чиройли CSS жадвалли HTML веб саҳифасига айлантиради"""
    rows_data = read_table_rows(input_path)
    output_filename = f"converted_{uuid.uuid4().hex[:8]}.html"
    output_path = str(CONVERT_DIR / output_filename)

    html_parts = [
        "<!DOCTYPE html>",
        "<html><head><meta charset='utf-8'><title>Жадвал</title>",
        "<style>",
        "body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; padding: 20px; background: #f4f6f9; }",
        "table { border-collapse: collapse; width: 100%; background: #fff; box-shadow: 0 1px 3px rgba(0,0,0,0.1); border-radius: 6px; overflow: hidden; }",
        "th, td { border: 1px solid #e2e8f0; padding: 8px 12px; text-align: left; font-size: 13px; }",
        "th { background-color: #1F4E79; color: white; font-weight: 600; }",
        "tr:nth-child(even) { background-color: #f7fafc; }",
        "tr:hover { background-color: #edf2f7; }",
        "</style></head><body>",
        "<table>",
    ]

    for idx, row in enumerate(rows_data):
        html_parts.append("<tr>")
        tag = "th" if idx == 0 else "td"
        for cell in row:
            html_parts.append(f"<{tag}>{html.escape(str(cell))}</{tag}>")
        html_parts.append("</tr>")

    html_parts.append("</table></body></html>")

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n".join(html_parts))

    return output_path


def convert_to_xlsx(input_path: str) -> str:
    """XLS ёки CSV файлни замонавий .xlsx Excel форматига ўтказади"""
    import openpyxl
    rows_data = read_table_rows(input_path)

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.xlsx"
    output_path = str(CONVERT_DIR / output_filename)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Жадвал"

    for row in rows_data:
        ws.append(row)

    wb.save(output_path)
    return output_path


# ==========================================
# РАСМЛАРНИ ҚАЙТА ИШЛАШ
# ==========================================
def convert_image(input_path: str, target_format: str) -> str:
    out_ext = target_format.lower()
    if out_ext == "jpg":
        out_ext = "jpeg"

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.{target_format.lower()}"
    output_path = str(CONVERT_DIR / output_filename)

    with Image.open(input_path) as img:
        if out_ext in ("jpeg", "bmp") and img.mode in ("RGBA", "LA", "P"):
            background = Image.new("RGB", img.size, (255, 255, 255))
            if img.mode == "P":
                img = img.convert("RGBA")
            background.paste(img, mask=img.split()[-1])
            background.save(output_path, format=out_ext.upper(), quality=95)
        elif out_ext == "ico":
            img.save(output_path, format="ICO", sizes=[(16, 16), (32, 32), (48, 48), (64, 64), (128, 128)])
        elif out_ext == "pdf":
            if img.mode != "RGB":
                img = img.convert("RGB")
            img.save(output_path, format="PDF", resolution=100.0)
        else:
            img.save(output_path)

    return output_path


# ==========================================
# АУДИО ВА ВИДЕО (FFmpeg)
# ==========================================
def find_ffmpeg() -> str:
    ffmpeg_cmd = shutil.which("ffmpeg")
    if ffmpeg_cmd:
        return ffmpeg_cmd

    # static_ffmpeg кутубхонаси орқали текшириш
    try:
        import static_ffmpeg
        static_ffmpeg.add_paths()
        ffmpeg_cmd = shutil.which("ffmpeg")
        if ffmpeg_cmd:
            return ffmpeg_cmd
    except Exception:
        pass

    # Linux стандарт йўллари
    for lp in ["/usr/bin/ffmpeg", "/usr/local/bin/ffmpeg", "/bin/ffmpeg"]:
        if os.path.exists(lp):
            return lp

    # Windows йўллари
    local_app_data = os.environ.get("LOCALAPPDATA", "")
    possible_paths = [
        Path(local_app_data) / "Microsoft" / "WinGet" / "Packages",
        Path("C:/Program Files/ffmpeg/bin"),
        Path("C:/ffmpeg/bin"),
    ]
    for p in possible_paths:
        if p.exists():
            found = list(p.glob("**/ffmpeg.exe"))
            if found:
                return str(found[0])

    return "ffmpeg"


def convert_media_with_ffmpeg(input_path: str, target_format: str) -> str:
    ffmpeg_exe = find_ffmpeg()
    target_format = target_format.lower()
    output_filename = f"converted_{uuid.uuid4().hex[:8]}.{target_format}"
    output_path = str(CONVERT_DIR / output_filename)

    cmd = [ffmpeg_exe, "-y", "-i", input_path]

    if target_format == "mp3":
        cmd.extend(["-vn", "-acodec", "libmp3lame", "-q:a", "2", output_path])
    elif target_format == "gif":
        cmd.extend(["-t", "15", "-vf", "fps=10,scale=480:-1:flags=lanczos", output_path])
    elif target_format == "wav":
        cmd.extend(["-vn", output_path])
    elif target_format == "ogg":
        cmd.extend(["-vn", "-acodec", "libopus", output_path])
    elif target_format == "m4a":
        cmd.extend(["-vn", "-c:a", "aac", output_path])
    elif target_format == "mp4":
        cmd.extend(["-c:v", "libx264", "-c:a", "aac", output_path])
    else:
        cmd.append(output_path)

    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"FFmpeg хатолиги: {res.stderr[:300]}")

    return output_path


# ==========================================
# ҲУЖЖАТЛАР (PDF, DOCX, TXT)
# ==========================================
def convert_pdf_to_docx(pdf_path: str) -> str:
    from pdf2docx import Converter

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.docx"
    output_path = str(CONVERT_DIR / output_filename)

    cv = Converter(pdf_path)
    cv.convert(output_path, start=0, end=None)
    cv.close()
    return output_path


def convert_pdf_to_txt(pdf_path: str) -> str:
    import fitz

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.txt"
    output_path = str(CONVERT_DIR / output_filename)

    doc = fitz.open(pdf_path)
    all_text = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        all_text.append(f"--- Саҳифа {page_num + 1} ---\n" + page.get_text())
    doc.close()

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(all_text))

    return output_path


def convert_pdf_to_images(pdf_path: str) -> str:
    import fitz
    import zipfile

    doc = fitz.open(pdf_path)
    total_pages = len(doc)

    if total_pages == 1:
        output_filename = f"converted_{uuid.uuid4().hex[:8]}.png"
        output_path = str(CONVERT_DIR / output_filename)
        page = doc[0]
        pix = page.get_pixmap(dpi=150)
        pix.save(output_path)
        doc.close()
        return output_path

    zip_filename = f"pdf_images_{uuid.uuid4().hex[:8]}.zip"
    zip_path = str(CONVERT_DIR / zip_filename)

    with zipfile.ZipFile(zip_path, "w") as zf:
        for i in range(total_pages):
            page = doc[i]
            pix = page.get_pixmap(dpi=150)
            img_bytes = pix.tobytes("png")
            zf.writestr(f"page_{i + 1}.png", img_bytes)

    doc.close()
    return zip_path


def convert_txt_to_pdf(txt_path: str) -> str:
    import fitz

    output_filename = f"converted_{uuid.uuid4().hex[:8]}.pdf"
    output_path = str(CONVERT_DIR / output_filename)

    with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
        text = f.read()

    doc = fitz.open()
    rect = fitz.Rect(50, 50, 550, 792)
    page = doc.new_page()
    page.insert_textbox(rect, text, fontsize=11, fontname="helv")
    doc.save(output_path)
    doc.close()

    return output_path


# ==========================================
# УМУМИЙ ДИСПЕТЧЕР
# ==========================================
def convert_file(input_path: str, target_format: str) -> str:
    _, ext = os.path.splitext(input_path)
    ext = ext.lower()
    target_format = target_format.lower()

    cat = get_file_category(ext)

    if cat == "excel":
        if target_format == "pdf":
            return convert_table_to_pdf(input_path)
        elif target_format == "csv":
            return convert_excel_to_csv(input_path)
        elif target_format == "html":
            return convert_excel_to_html(input_path)
        elif target_format == "xlsx":
            return convert_to_xlsx(input_path)

    if cat == "image":
        return convert_image(input_path, target_format)

    if cat in ("audio", "video"):
        return convert_media_with_ffmpeg(input_path, target_format)

    if cat == "document":
        if ext == ".pdf":
            if target_format == "docx":
                return convert_pdf_to_docx(input_path)
            elif target_format == "txt":
                return convert_pdf_to_txt(input_path)
            elif target_format == "png_images":
                return convert_pdf_to_images(input_path)
        elif ext == ".txt" and target_format == "pdf":
            return convert_txt_to_pdf(input_path)

    raise ValueError(f"{ext} форматидан {target_format} форматига ўгириб бўлмайди.")
